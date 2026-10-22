"""Excel results workbook (services/irs_outputs/excel_report.py) for all ten cases.

Each case is priced from its worksheet in the v4 input workbook, written to a results
workbook, read back, and checked against the engine's outputs and the approved Optima
regression baseline (tests/fixtures/irs_expected_outputs, suite tolerances).
"""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts import irs_expected_output_suite as suite
from scripts import price_irs_workbook as cli
from services.irs_inputs.pricing import price_workbook
from services.irs_outputs.excel_report import ReportExistsError, write_pricing_report

ROOT = Path(__file__).resolve().parents[3]
WORKBOOK = ROOT / "examples" / "irs" / "Optima_IRS_Test_Case_Inputs_v4.xlsx"
SHEETS = {f"TC{n:03d}": f"TC{n}" for n in range(1, 11)}
SHEET_NAMES = ["Summary", "Fixed Leg", "Floating Leg", "Net Cashflow"]
TOL = suite.TOLERANCES


def _day(value):
    return value.date() if isinstance(value, datetime) else value


def _table(ws):
    rows = list(ws.iter_rows(values_only=True))
    header, body = rows[0], rows[1:]
    data = [dict(zip(header, r)) for r in body if r[0] is not None]
    total = [r for r in body if r[0] is None]
    return data, total[0] if total else None


def _summary(ws):
    return {r[0]: r[1] for r in ws.iter_rows(min_row=3, values_only=True) if r[0]}


@pytest.fixture(scope="module")
def reports(tmp_path_factory):
    folder = tmp_path_factory.mktemp("results")
    out = {}
    for sheet, case in SHEETS.items():
        priced = price_workbook(WORKBOOK, sheet)
        path = write_pricing_report(priced, folder / f"{sheet}.xlsx",
                                    source=f"{WORKBOOK.name} [{sheet}]")
        out[sheet] = (priced, load_workbook(path), json.loads(
            suite.expected_path(case).read_text(encoding="utf-8"))["expected"])
    return out


@pytest.mark.parametrize("sheet", sorted(SHEETS))
def test_sheets_and_summary(reports, sheet):
    priced, wb, expected = reports[sheet]
    assert wb.sheetnames == SHEET_NAMES
    s = _summary(wb["Summary"])
    assert (s["Trade ID"], s["Currency"], _day(s["Valuation Date"])) == (
        priced.trade_id, priced.currency, priced.valuation_date)
    # Excel stores numbers to about 15 significant digits.
    assert s["Swap NPV"] == pytest.approx(priced.npv, abs=1e-9)
    assert s["Swap NPV"] == pytest.approx(expected["npv"], abs=TOL["currency"])
    assert s["Fixed Leg NPV"] == pytest.approx(expected["fixed_leg_npv"], abs=TOL["currency"])
    assert s["Floating Leg NPV"] == pytest.approx(expected["floating_leg_npv"],
                                                  abs=TOL["currency"])
    assert s["Market Data"].split("\n") == [p.name for p in priced.market_files]
    if sheet == "TC002":
        assert "PROVISIONAL pending SME decision R9" in s["Notes"]
    else:
        assert s["Notes"] == "-"


@pytest.mark.parametrize("sheet", sorted(SHEETS))
@pytest.mark.parametrize("leg, title", [("fixed", "Fixed Leg"), ("floating", "Floating Leg")])
def test_leg_sheets_match_the_baseline(reports, sheet, leg, title):
    priced, wb, expected = reports[sheet]
    rows, total = _table(wb[title])
    instrument = priced.engine_outputs[0]
    periods = (instrument.trade.fixed_periods if leg == "fixed"
               else instrument.trade.floating_periods)
    assert [r["Period"] for r in rows] == list(range(1, len(periods) + 1))
    notionals = [float(p.notional.amount) for p in periods]
    assert [r["Notional"] for r in rows] == notionals
    assert [r["Amortisation"] for r in rows] == [0.0] + [
        a - b for a, b in zip(notionals, notionals[1:])]
    baseline = {(c["leg"], c["period"]): c for c in expected["cashflows"]}
    settled = {(s["leg"], s["period"]) for s in expected["settled"]}
    projected = [r for r in rows if r["Status"] == "Projected"]
    for r in rows:
        key = (leg, r["Period"])
        if r["Status"] == "Settled":
            assert key in settled and r["Amount"] is None and r["Present Value"] is None
            continue
        c = baseline[key]
        assert (_day(r["Accrual Start"]).isoformat(), _day(r["Accrual End"]).isoformat(),
                _day(r["Payment Date"]).isoformat()) == (
            c["accrual_start"], c["accrual_end"], c["payment_date"])
        rate = r["Rate incl. Spread"] if leg == "floating" else r["Fixed Rate"]
        assert rate == pytest.approx(c["rate"], abs=TOL["rate"])
        assert r["Day Count Fraction"] == pytest.approx(c["day_count_fraction"],
                                                        abs=TOL["day_count_fraction"])
        assert r["Amount"] == pytest.approx(c["amount"], abs=TOL["currency"])
        assert r["Discount Factor"] == pytest.approx(c["discount_factor"],
                                                     abs=TOL["discount_factor"])
        assert r["Present Value"] == pytest.approx(c["present_value"], abs=TOL["currency"])
    assert len(projected) == len([k for k in baseline if k[0] == leg])
    leg_npv = expected[f"{leg}_leg_npv"]
    assert total[-1] == pytest.approx(leg_npv, abs=TOL["currency"])
    assert sum(r["Present Value"] for r in projected) == pytest.approx(leg_npv, abs=1e-6)


@pytest.mark.parametrize("sheet", sorted(SHEETS))
def test_net_cashflow_per_payment_date(reports, sheet):
    priced, wb, expected = reports[sheet]
    net, total = _table(wb["Net Cashflow"])
    fixed, _ = _table(wb["Fixed Leg"])
    floating, _ = _table(wb["Floating Leg"])
    dates = sorted({_day(r["Payment Date"]) for r in fixed + floating})
    assert [_day(r["Payment Date"]) for r in net] == dates
    for r in net:
        if r["Status"] == "Settled":
            assert r["Net Amount"] is None
            continue
        assert r["Net Amount"] == pytest.approx(r["Fixed Amount"] + r["Floating Amount"],
                                                abs=1e-9)
        assert r["Net Present Value"] == pytest.approx(r["Net Amount"] * r["Discount Factor"],
                                                       abs=1e-6)
    assert total[-1] == pytest.approx(priced.npv, abs=1e-9)
    assert sum(r["Net Present Value"] for r in net if r["Status"] == "Projected") == \
        pytest.approx(expected["npv"], abs=1e-6)
    # amortisation: initial notional = what is outstanding at the end + all reductions
    assert net[0]["Notional"] == pytest.approx(
        net[-1]["Notional"] + sum(r["Amortisation"] for r in net[1:]), abs=1e-9)


def test_settled_periods_are_listed_not_valued(reports):
    """TC4 has one settled period per leg (paid on the valuation date)."""
    _, wb, _ = reports["TC004"]
    net, _ = _table(wb["Net Cashflow"])
    assert [(_day(r["Payment Date"]), r["Status"]) for r in net][0] == (
        date(2026, 3, 31), "Settled")


def test_results_hold_values_only(reports):
    for _, wb, _ in reports.values():
        for ws in wb.worksheets:
            assert not any(isinstance(c.value, str) and c.value.startswith("=")
                           for row in ws.iter_rows() for c in row)


def test_existing_file_is_not_overwritten(tmp_path):
    priced = price_workbook(WORKBOOK, "TC001")
    path = write_pricing_report(priced, tmp_path / "r.xlsx")
    before = path.read_bytes()
    with pytest.raises(ReportExistsError):
        write_pricing_report(priced, path)
    assert path.read_bytes() == before
    write_pricing_report(priced, path, overwrite=True)
    with pytest.raises(ValueError, match=".xlsx"):
        write_pricing_report(priced, tmp_path / "r.csv")


def test_cli_saves_results(tmp_path, capsys):
    out = tmp_path / "TC005.xlsx"
    assert cli.main(["--workbook", str(WORKBOOK), "--sheet", "TC005", "--output", str(out)]) == 0
    assert "Results saved" in capsys.readouterr().out
    assert load_workbook(out).sheetnames == SHEET_NAMES
    assert cli.main(["--workbook", str(WORKBOOK), "--sheet", "TC005", "--output", str(out)]) == 2
    assert "not overwritten" in capsys.readouterr().err
    assert cli.main(["--workbook", str(WORKBOOK), "--sheet", "TC005", "--output", str(out),
                     "--overwrite"]) == 0
