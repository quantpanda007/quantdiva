"""TC1 from Excel to the existing engine: reader -> normalizer -> pricing orchestration.

Inputs: examples/irs/Optima_IRS_Test_Case_Inputs_v4.xlsx (worksheet TC001), populated from
examples/irs/tc1_trade.json. Expected output: the approved Optima regression baseline
tests/fixtures/irs_expected_outputs/TC1.json (kept separate from the workbook; not altered),
compared with the suite's own tolerances.
"""

from __future__ import annotations

import ast
import copy
import dataclasses
import json
import shutil
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
import QuantLib as ql
from openpyxl import load_workbook

import price_irs_engine as engine
from instruments.rates.irs import floating_leg as engine_floating
from instruments.rates.irs import pricing_instrument as engine_schema
from scripts import irs_expected_output_suite as suite
from scripts import price_irs_workbook as cli
from services.irs_inputs import normalizer
from services.irs_inputs.contract import (
    SCHEDULE_ROW_KEYS,
    SECTION_KEYS,
    InputValidationError,
    NormalizedIRSTrade,
)
from services.irs_inputs.excel_reader import WorkbookReadError, read_trade_sheet
from services.irs_inputs.normalizer import normalize
from services.irs_inputs.pricing import price_normalized, price_workbook

ROOT = Path(__file__).resolve().parents[3]
WORKBOOK = ROOT / "examples" / "irs" / "Optima_IRS_Test_Case_Inputs_v4.xlsx"
SHEET = "TC001"
TC1_TRADE = ROOT / "examples" / "irs" / "tc1_trade.json"
BASELINE = ROOT / "tests" / "fixtures" / "irs_expected_outputs" / "TC1.json"
USD_SNAPSHOT = ROOT / "market_data" / "2026-03-31" / "USD" / "usd_sofr_2026-03-31.json"
PACKAGE = ROOT / "services" / "irs_inputs"

# What the TC1 worksheet holds (analyst units and labels), read cell by cell.
TC1_INPUTS = {
    "trade": {"trade_id": "SLLSL9RR", "currency": "USD", "notional": 530000,
              "effective_date": date(2026, 5, 31), "maturity_date": date(2027, 11, 30),
              "valuation_date": date(2026, 3, 31)},
    "fixed_leg": {"direction": "Pay", "rate_percent": Decimal("5.00"),
                  "payment_frequency": "Annual", "day_count": "ACT/360",
                  "business_day_convention": "Modified Following", "calendar": "FD"},
    "floating_leg": {"direction": "Receive", "benchmark": "SOFR", "index_tenor": "1D",
                     "spread_bp": 10, "leverage": 1, "reset_frequency": "Daily",
                     "payment_frequency": "Annual", "day_count": "ACT/360",
                     "business_day_convention": "Modified Following", "calendar": "FD",
                     "rate_timing": "In Arrears", "lookback_days": 0, "lookback_type": "Lookback",
                     "rate_cutoff_days": 1, "compounding": "Daily Compounded",
                     "spread_treatment": "Exclude Spread"},
    "notional": {"principal_exchange": "Never", "amortisation_type": "Custom Schedule",
                 "schedule": [
                     {"period": 1, "accrual_start": date(2026, 5, 31),
                      "accrual_end": date(2027, 5, 28), "payment_date": date(2027, 5, 28),
                      "notional": 530000, "amortisation_percent": None,
                      "amortisation_amount": 0, "action": None},
                     {"period": 2, "accrual_start": date(2027, 5, 28),
                      "accrual_end": date(2027, 11, 30), "payment_date": date(2027, 11, 30),
                      "notional": 265000, "amortisation_percent": None,
                      "amortisation_amount": 265000, "action": None},
                 ]},
    "market_data": {"curve": "USD SOFR", "snapshot_folder": "market_data/2026-03-31/USD/",
                    "curve_source_file": "usd_sofr_2026-03-31.json",
                    "fixings_source_file": None, "reference_date": None,
                    "interpolation": "Piecewise Linear (Cont)"},
}


def _json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def inputs():
    return read_trade_sheet(WORKBOOK, SHEET)


@pytest.fixture(scope="module")
def normalized(inputs):
    return normalize(inputs)


def _mutated(inputs, section, key, value):
    doc = copy.deepcopy(inputs)
    doc[section][key] = value
    return doc


def _issues(doc) -> str:
    with pytest.raises(InputValidationError) as info:
        normalize(doc)
    return str(info.value)


# --- 1. the workbook is read correctly ---------------------------------------------------------


def test_tc1_worksheet_is_read_cell_by_cell(inputs):
    assert inputs == TC1_INPUTS


def test_sheet_must_be_named_unless_the_workbook_has_only_one(inputs, tmp_path):
    with pytest.raises(WorkbookReadError, match="name the one to read"):
        read_trade_sheet(WORKBOOK)
    single = tmp_path / "single.xlsx"
    shutil.copy(WORKBOOK, single)
    wb = load_workbook(single)
    for name in wb.sheetnames:
        if name != SHEET:
            del wb[name]
    wb.save(single)
    single_inputs, expected = read_trade_sheet(single), copy.deepcopy(inputs)
    # A workbook saved by openpyxl has no calculated value for the Snapshot Folder formula.
    assert single_inputs["market_data"].pop("snapshot_folder") is None
    expected["market_data"].pop("snapshot_folder")
    assert single_inputs == expected


def test_reader_output_follows_the_contract(inputs):
    assert set(inputs) == set(SECTION_KEYS)
    for section, keys in SECTION_KEYS.items():
        assert set(inputs[section]) == set(keys)
    assert all(set(row) == set(SCHEDULE_ROW_KEYS) for row in inputs["notional"]["schedule"])


# --- 2. mapping into the engine representation -------------------------------------------------


def test_normalized_trade_is_exactly_the_existing_tc1_fixture(normalized):
    """Field by field, units and enumerations included: the Excel route reproduces the
    authoritative TC1 trade input."""
    assert normalized.trade_document == _json(TC1_TRADE)
    assert json.dumps(normalized.trade_document) == json.dumps(_json(TC1_TRADE))


def test_market_request(normalized):
    m = normalized.market
    assert (m.valuation_date, m.currency, m.curve_benchmark, m.curve_file, m.fixings_file,
            m.reference_date, m.interpolation) == (
        date(2026, 3, 31), "USD", "SOFR", "usd_sofr_2026-03-31.json", None, None, "linear_zero")
    assert normalized.notes == ()


def test_normalized_representation_has_no_excel_dependency():
    for name in ("contract.py", "normalizer.py", "pricing.py"):
        tree = ast.parse((PACKAGE / name).read_text(encoding="utf-8"))
        modules = {n.module for n in ast.walk(tree) if isinstance(n, ast.ImportFrom)} | {
            a.name for n in ast.walk(tree) if isinstance(n, ast.Import) for a in n.names}
        assert not any(m and m.startswith("openpyxl") for m in modules), name
        if name != "pricing.py":  # pricing offers price_workbook as a convenience only
            assert "services.irs_inputs.excel_reader" not in modules, name


def test_a_ui_style_payload_gives_the_same_trade(inputs, normalized):
    """JSON-style values (ISO dates, numeric strings, floats) normalize identically."""
    payload = json.loads(json.dumps(inputs, default=str))
    payload["fixed_leg"]["rate_percent"] = 5.0
    assert normalize(payload).trade_document == normalized.trade_document


# --- 3. units and dates ------------------------------------------------------------------------


@pytest.mark.parametrize("percent, rate", [(Decimal("5.00"), 0.05), ("4.25", 0.0425), (3, 0.03)])
def test_rate_percent_becomes_decimal(inputs, percent, rate):
    doc = _mutated(inputs, "fixed_leg", "rate_percent", percent)
    assert normalize(doc).trade_document["params"]["fixed_leg"]["fixed_rate"] == rate


@pytest.mark.parametrize("bp, spread", [(10, 0.001), (20, 0.002), ("2.5", 0.00025), (0, 0.0)])
def test_spread_bp_becomes_decimal(inputs, bp, spread):
    doc = _mutated(inputs, "floating_leg", "spread_bp", bp)
    assert normalize(doc).trade_document["params"]["floating_leg"]["spread"] == spread


def test_dates_accept_date_datetime_and_iso_text(inputs, normalized):
    doc = copy.deepcopy(inputs)
    doc["trade"]["effective_date"] = datetime(2026, 5, 31)
    doc["trade"]["maturity_date"] = "2027-11-30"
    assert normalize(doc).trade_document == normalized.trade_document


@pytest.mark.parametrize("value, message", [
    (datetime(2026, 5, 31, 12, 0), "without a time of day"),
    ("31/05/2026", "must be a date"),
    (46173, "must be a date"),
])
def test_bad_dates_are_rejected(inputs, value, message):
    assert message in _issues(_mutated(inputs, "trade", "effective_date", value))


def _workbook_copy(tmp_path: Path) -> Path:
    target = tmp_path / "inputs.xlsx"
    shutil.copy(WORKBOOK, target)
    return target


def _edit(path: Path, **cells):
    wb = load_workbook(path)
    ws = wb[SHEET]
    for address, value in cells.items():
        ws[address] = value
    wb.save(path)


def _find_row(label: str, after: str = None) -> int:
    ws = load_workbook(WORKBOOK)[SHEET]
    started = after is None
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if value == after:
            started = True
        elif started and value == label:
            return row
    raise AssertionError(label)


def test_fixed_rate_must_be_a_percentage_cell(tmp_path):
    path = _workbook_copy(tmp_path)
    row = _find_row("Fixed Rate")
    wb = load_workbook(path)
    wb[SHEET][f"B{row}"].number_format = "0.00"
    wb.save(path)
    with pytest.raises(WorkbookReadError, match="must be a percentage-formatted cell"):
        read_trade_sheet(path, SHEET)


def test_amortisation_percent_cell_is_read_as_percent(tmp_path):
    path = _workbook_copy(tmp_path)
    period_1 = _find_row("Period") + 1
    _edit(path, **{f"F{period_1}": 0.0})
    row = read_trade_sheet(path, SHEET)["notional"]["schedule"][0]
    assert row["amortisation_percent"] == 0
    _edit(path, **{f"F{period_1 + 1}": 0.5})
    assert read_trade_sheet(path, SHEET)["notional"]["schedule"][1][
        "amortisation_percent"] == Decimal("50.0")


# --- 4. missing, invalid and unsupported inputs ---------------------------------------------


@pytest.mark.parametrize("section, key, value, message", [
    ("trade", "trade_id", None, "trade.trade_id: is required"),
    ("trade", "currency", "EUR", "'EUR' is not a recognised value"),
    ("trade", "notional", "abc", "trade.notional: must be a number"),
    ("fixed_leg", "rate_percent", None, "fixed_leg.rate_percent: is required"),
    ("fixed_leg", "calendar", "XX", "fixed_leg.calendar: 'XX' is not a recognised value"),
    ("fixed_leg", "payment_frequency", "Weekly", "'Weekly' is not a recognised value"),
    ("floating_leg", "direction", "Pay", "must be the opposite of the fixed leg"),
    ("floating_leg", "leverage", None, "floating_leg.leverage: is required"),
    ("floating_leg", "calendar", None, "floating_leg.calendar: is required"),
    ("floating_leg", "business_day_convention", None,
     "floating_leg.business_day_convention: is required"),
    ("floating_leg", "lookback_type", "Observation Shift", "not a recognised value"),
    ("floating_leg", "lookback_days", 1.5, "must be a whole number of days"),
    ("floating_leg", "rate_cutoff_days", -1, "must be a whole number of days"),
    ("floating_leg", "benchmark", "SARON", "is a term-rate (IBOR-fallback) leg"),
    ("floating_leg", "benchmark", "INR OIS", "is a term-rate (IBOR-fallback) leg"),
    ("floating_leg", "benchmark", "TONAR", "TONAR 1D is not a supported USD overnight"),
    ("floating_leg", "index_tenor", "6M", "is a term-rate (IBOR-fallback) leg"),
    ("floating_leg", "reset_frequency", "Monthly", "only daily reset"),
    ("floating_leg", "compounding", "Simple", "'Simple' not supported by the engine"),
    ("floating_leg", "spread_treatment", "Compound Spread", "not supported by the engine"),
    ("notional", "principal_exchange", "At End", "principal exchange is not supported"),
    ("notional", "amortisation_type", "Percentage per Period", "semantics not yet defined"),
    ("market_data", "interpolation", "Linear", "meaning not yet defined"),
    ("market_data", "interpolation", "Log Linear", "meaning not yet defined"),
    ("market_data", "curve", "JPY TONAR", "is not a USD curve"),
    ("market_data", "curve_source_file", None, "market_data.curve_source_file: is required"),
    ("market_data", "curve_source_file", "usd_sofr_2026-03-30.json", "is not the prepared"),
    ("market_data", "snapshot_folder", "market_data/2026-03-31/CHF/", "does not match"),
])
def test_invalid_input_gives_a_clear_error(inputs, section, key, value, message):
    assert message in _issues(_mutated(inputs, section, key, value))


@pytest.mark.parametrize("row, key, value, message", [
    (1, "accrual_start", date(2026, 6, 1), "must equal the effective date"),
    (2, "accrual_start", date(2027, 5, 31), "periods must be contiguous"),
    (1, "notional", 500000, "must equal trade.notional"),
    (2, "amortisation_amount", 200000, "is inconsistent with the notionals"),
    (2, "amortisation_percent", 50, "its base (original or outstanding notional) is not yet"),
    (1, "action", "Amortise", "has no defined meaning yet"),
    (2, "period", 3, "numbered 1, 2, ... in order"),
    (2, "notional", 0, "must be positive"),
])
def test_invalid_schedule_gives_a_clear_error(inputs, row, key, value, message):
    doc = copy.deepcopy(inputs)
    doc["notional"]["schedule"][row - 1][key] = value
    assert message in _issues(doc)


def test_zero_amortisation_percent_is_unambiguous_and_accepted(inputs, normalized):
    doc = copy.deepcopy(inputs)
    doc["notional"]["schedule"][0]["amortisation_percent"] = 0
    assert normalize(doc).trade_document == normalized.trade_document


def test_bullet_requires_constant_notional(inputs):
    doc = _mutated(inputs, "notional", "amortisation_type", "Bullet / No Amortisation")
    assert "requires every period notional to equal 530000" in _issues(doc)


def test_every_problem_is_reported_at_once(inputs):
    doc = copy.deepcopy(inputs)
    doc["trade"]["trade_id"] = None
    doc["floating_leg"]["leverage"] = None
    doc["market_data"]["interpolation"] = "Linear"
    with pytest.raises(InputValidationError) as info:
        normalize(doc)
    assert len(info.value.issues) == 3


def test_missing_and_unknown_fields_are_rejected(inputs):
    doc = copy.deepcopy(inputs)
    del doc["floating_leg"]["leverage"]
    doc["fixed_leg"]["payment_lag"] = 2
    text = _issues(doc)
    assert "floating_leg.leverage: field is missing" in text
    assert "fixed_leg.payment_lag: is not a field of the input contract" in text


def test_engine_rejection_is_reported_as_an_input_error(inputs):
    doc = _mutated(inputs, "trade", "maturity_date", date(2027, 12, 31))
    assert "engine rejected the trade" in _issues(doc)


def test_v3_layout_without_the_new_rows_is_rejected_with_the_reason(tmp_path):
    """The v3 layout lacks the four floating-leg rows; nothing is defaulted."""
    path = _workbook_copy(tmp_path)
    wb = load_workbook(path)
    ws = wb[SHEET]
    first = _find_row("Leverage")
    rows = [first, _find_row("Business Day Convention", after="SWAP LEGS — FLOATING LEG"),
            _find_row("Calendar"), _find_row("Lookback Type")]
    for row in sorted(rows, reverse=True):
        ws.delete_rows(row)
    wb.save(path)
    with pytest.raises(WorkbookReadError, match="has no rows .*Leverage"):
        read_trade_sheet(path, SHEET)


@pytest.mark.parametrize("edit, message", [
    (lambda ws: ws.cell(_find_row("Leverage"), 1, "Gearing"), "unknown label 'Gearing'"),
    (lambda ws: setattr(ws.cell(_find_row("Lookback Type"), 1), "value", None),
     "value without a label"),
    (lambda ws: ws.cell(_find_row("Period") + 4, 2, datetime(2028, 1, 1)),
     "amortisation row after an empty row"),
])
def test_unexpected_layout_is_rejected(tmp_path, edit, message):
    path = _workbook_copy(tmp_path)
    wb = load_workbook(path)
    edit(wb[SHEET])
    wb.save(path)
    with pytest.raises(WorkbookReadError, match=message):
        read_trade_sheet(path, SHEET)


def test_sheet_selection_errors(tmp_path):
    path = _workbook_copy(tmp_path)
    with pytest.raises(WorkbookReadError, match="no worksheet 'TC002'"):
        read_trade_sheet(path, "TC002")
    wb = load_workbook(path)
    wb.copy_worksheet(wb[SHEET]).title = "TC002"
    wb.save(path)
    with pytest.raises(WorkbookReadError, match="name the one to read"):
        read_trade_sheet(path)


def test_not_a_workbook_is_rejected(tmp_path):
    bogus = tmp_path / "bogus.xlsx"
    bogus.write_text("not excel", encoding="utf-8")
    with pytest.raises(WorkbookReadError, match="cannot open as an Excel workbook"):
        read_trade_sheet(bogus)


# --- 5. market-data resolution -------------------------------------------------------------


def test_market_data_resolves_to_the_canonical_snapshot_only():
    priced = price_workbook(WORKBOOK, SHEET)
    assert priced.market_files == (USD_SNAPSHOT,)
    assert priced.fixing_reasons == ()


def _with_market(normalized, **changes) -> NormalizedIRSTrade:
    return dataclasses.replace(normalized,
                               market=dataclasses.replace(normalized.market, **changes))


@pytest.mark.parametrize("changes, message", [
    ({"fixings_file": "fixings_2026-03-31.json"}, "needs no fixing"),
    ({"reference_date": date(2026, 3, 30)}, "prepared snapshot USD-SOFR-490 has 2026-03-31"),
    ({"interpolation": "log_linear_discount"}, "uses 'linear_zero'"),
])
def test_market_reference_must_agree_with_the_snapshot(normalized, changes, message):
    with pytest.raises(InputValidationError, match=message):
        price_normalized(_with_market(normalized, **changes))


def test_stated_reference_date_equal_to_snapshot_is_accepted(normalized):
    priced = price_normalized(_with_market(normalized, reference_date=date(2026, 3, 31)))
    assert priced.npv == pytest.approx(_json(BASELINE)["expected"]["npv"], abs=1e-6)


def test_a_trade_needing_fixings_must_name_the_dated_file(normalized):
    """TC4-like: a SOFR period running at the valuation date needs the shared fixings."""
    doc = copy.deepcopy(_json(ROOT / "examples" / "irs" / "tc4_trade.json"))
    trade = dataclasses.replace(normalized, trade_document=doc,
                                market=dataclasses.replace(normalized.market,
                                                           fixings_file=None))
    with pytest.raises(InputValidationError, match="needs fixings"):
        price_normalized(trade)
    named = _with_market(trade, fixings_file="fixings_2026-03-31.json")
    assert price_normalized(named).market_files[1].name == "fixings_2026-03-31.json"


def test_missing_market_root_fails_clearly(normalized, tmp_path):
    with pytest.raises(engine.MarketDataResolutionError, match="no USD SOFR market snapshot"):
        price_normalized(normalized, tmp_path)


def test_market_files_are_not_modified_by_pricing():
    before = USD_SNAPSHOT.read_bytes()
    price_workbook(WORKBOOK, SHEET)
    assert USD_SNAPSHOT.read_bytes() == before


# --- 6. the normalized trade reaches the existing engine ------------------------------------


def test_normalized_trade_is_what_the_engine_prices(normalized, monkeypatch):
    seen = []
    real = engine.price

    def recording(trade_doc, market_doc):
        seen.append(trade_doc)
        return real(trade_doc, market_doc)

    monkeypatch.setattr(engine, "price", recording)
    price_normalized(normalized)
    assert seen == [normalized.trade_document]


def test_adapter_contains_no_pricing_formula():
    """Orchestration only: no QuantLib, discounting or cashflow arithmetic of its own."""
    for name in ("contract.py", "normalizer.py", "pricing.py", "excel_reader.py"):
        text = (PACKAGE / name).read_text(encoding="utf-8")
        assert "import QuantLib" not in text and "discount_factor *" not in text, name


# --- 7. end-to-end TC1 against the approved baseline -----------------------------------------


def _suite_outputs(priced) -> dict:
    return {
        "npv": priced.npv, "fixed_leg_npv": priced.fixed_leg_npv,
        "floating_leg_npv": priced.floating_leg_npv,
        "cashflows": [
            {"leg": r.leg, "side": r.side, "period": r.period,
             "accrual_start": r.accrual_start.isoformat(),
             "accrual_end": r.accrual_end.isoformat(),
             "payment_date": r.payment_date.isoformat(), "notional": r.notional,
             "day_count_fraction": r.day_count_fraction, "rate": r.rate, "amount": r.amount,
             "discount_factor": r.discount_factor, "present_value": r.present_value}
            for r in priced.cashflows],
        "settled": [{"leg": s.leg, "side": s.side, "period": s.period,
                     "payment_date": s.payment_date.isoformat()} for s in priced.settled],
    }


def test_tc1_from_excel_matches_the_approved_baseline():
    baseline = _json(BASELINE)
    assert baseline["provenance"]["kind"] == "optima_regression_baseline"
    priced = price_workbook(WORKBOOK, SHEET)
    assert priced.npv == pytest.approx(-7955.726204, abs=suite.TOLERANCES["currency"])
    result = suite.CaseResult("TC1", "PASS")
    suite.compare(_suite_outputs(priced), baseline["expected"], result)
    assert result.failures == []
    assert result.fields_compared == ["npv", "fixed_leg_npv", "floating_leg_npv", "cashflows",
                                      "settled"]


def test_tc1_pricing_is_reproducible_and_restores_the_evaluation_date():
    before = ql.Settings.instance().evaluationDate
    first, second = price_workbook(WORKBOOK, SHEET), price_workbook(WORKBOOK, SHEET)
    assert first.npv == second.npv
    assert [r.amount for r in first.cashflows] == [r.amount for r in second.cashflows]
    assert ql.Settings.instance().evaluationDate == before


def test_excel_route_equals_the_json_route():
    """Same engine, same market data: the Excel route and the existing trade-file route give
    the same outputs."""
    trade_doc = _json(TC1_TRADE)
    selection = engine.resolve_market_files(trade_doc)
    _, _, result, legs, rows, _ = engine.price(trade_doc,
                                               engine.merge_market(list(selection.files)))
    priced = price_workbook(WORKBOOK, SHEET)
    assert (priced.npv, priced.fixed_leg_npv, priced.floating_leg_npv) == (
        result.npv, legs["fixed"], legs["floating"])
    assert [r.amount for r in priced.cashflows] == [r.amount for r in rows]


def test_workbook_holds_no_expected_results():
    """The expected output stays in the baseline file, never in the input workbook."""
    for ws in load_workbook(WORKBOOK, data_only=True).worksheets:
        labels = {str(c.value).lower() for row in ws.iter_rows() for c in row if c.value}
        assert not any(word in label for label in labels for word in ("npv", "expected", "pv "))


def test_cli(capsys):
    assert cli.main(["--workbook", str(WORKBOOK), "--sheet", SHEET]) == 0
    out = capsys.readouterr().out
    assert "NPV -7,955.7262" in out and "usd_sofr_2026-03-31.json" in out


def test_cli_reports_input_errors(tmp_path, capsys):
    path = _workbook_copy(tmp_path)
    _edit(path, **{f"B{_find_row('Leverage')}": None})
    assert cli.main(["--workbook", str(path), "--sheet", SHEET]) == 2
    assert "floating_leg.leverage: is required" in capsys.readouterr().err


# --- 8. the input contract stays in step with the engine and the workbook ---------------------


def test_supported_values_match_the_engine():
    assert sorted(normalizer.CALENDARS) == sorted(engine_schema._CALENDARS)
    assert set(normalizer.DAY_COUNTS.values()) == set(engine_schema._DAY_COUNTS)
    assert set(normalizer.PAYMENT_FREQUENCIES.values()) == set(
        engine_schema._PAYMENT_FREQUENCIES)
    assert set(normalizer.RESET_FREQUENCIES.values()) == set(engine_schema._RESET_FREQUENCIES)
    assert set(normalizer.BUSINESS_DAY_CONVENTIONS.values()) <= set(
        engine_schema._BUSINESS_DAY_CONVENTIONS)
    engine_overnight = {(b, t, c) for b, t, c, _ in engine_floating._OVERNIGHT_BENCHMARKS}
    assert {(code, tenor, ccy) for (_, tenor, ccy), code in
            normalizer.OVERNIGHT_BENCHMARKS.items()} == engine_overnight


# Workbook field (column A label of the list cell) -> contract path.
_LIST_FIELDS = {
    ("TRADE DETAILS", "Currency"): "trade.currency",
    ("SWAP LEGS — FIXED LEG", "Direction"): "fixed_leg.direction",
    ("SWAP LEGS — FIXED LEG", "Payment Frequency"): "fixed_leg.payment_frequency",
    ("SWAP LEGS — FIXED LEG", "Day Count"): "fixed_leg.day_count",
    ("SWAP LEGS — FIXED LEG", "Business Day Convention"): "fixed_leg.business_day_convention",
    ("SWAP LEGS — FLOATING LEG", "Direction"): "floating_leg.direction",
    ("SWAP LEGS — FLOATING LEG", "Benchmark"): "floating_leg.benchmark",
    ("SWAP LEGS — FLOATING LEG", "Index Tenor"): "floating_leg.index_tenor",
    ("SWAP LEGS — FLOATING LEG", "Reset Frequency"): "floating_leg.reset_frequency",
    ("SWAP LEGS — FLOATING LEG", "Payment Frequency"): "floating_leg.payment_frequency",
    ("SWAP LEGS — FLOATING LEG", "Day Count"): "floating_leg.day_count",
    ("SWAP LEGS — FLOATING LEG", "Business Day Convention"):
        "floating_leg.business_day_convention",
    ("SWAP LEGS — FLOATING LEG", "Rate Timing"): "floating_leg.rate_timing",
    ("SWAP LEGS — FLOATING LEG", "Lookback Type"): "floating_leg.lookback_type",
    ("SWAP LEGS — FLOATING LEG", "Compounding"): "floating_leg.compounding",
    ("SWAP LEGS — FLOATING LEG", "Spread Treatment"): "floating_leg.spread_treatment",
    ("NOTIONAL & AMORTISATION", "Principal Exchange"): "notional.principal_exchange",
    ("NOTIONAL & AMORTISATION", "Amortisation Type"): "notional.amortisation_type",
    ("MARKET DATA REFERENCE", "Curve"): "market_data.curve",
    ("MARKET DATA REFERENCE", "Interpolation"): "market_data.interpolation",
}
_SUPPORTED = {
    "trade.currency": normalizer.CURRENCIES,
    "fixed_leg.direction": normalizer.DIRECTIONS, "floating_leg.direction": normalizer.DIRECTIONS,
    "fixed_leg.payment_frequency": normalizer.PAYMENT_FREQUENCIES,
    "floating_leg.payment_frequency": normalizer.PAYMENT_FREQUENCIES,
    "fixed_leg.day_count": normalizer.DAY_COUNTS, "floating_leg.day_count": normalizer.DAY_COUNTS,
    "fixed_leg.business_day_convention": normalizer.BUSINESS_DAY_CONVENTIONS,
    "floating_leg.business_day_convention": normalizer.BUSINESS_DAY_CONVENTIONS,
    "floating_leg.benchmark": {b for b, _, _ in normalizer.OVERNIGHT_BENCHMARKS},
    "floating_leg.index_tenor": {"1D"},
    "floating_leg.reset_frequency": normalizer.RESET_FREQUENCIES,
    "floating_leg.rate_timing": normalizer.RATE_TIMINGS,
    "floating_leg.lookback_type": normalizer.LOOKBACK_TYPES,
    "floating_leg.compounding": normalizer.COMPOUNDING,
    "floating_leg.spread_treatment": normalizer.SPREAD_TREATMENTS,
    "notional.principal_exchange": normalizer.PRINCIPAL_EXCHANGES,
    "notional.amortisation_type": normalizer.AMORTISATION_TYPES,
    "market_data.curve": normalizer.CURVES,
    "market_data.interpolation": normalizer.INTERPOLATIONS,
}


def test_every_workbook_option_is_supported_or_rejected_with_a_reason():
    """No workbook list option is silently mapped or silently unknown: each is supported,
    not supported by the engine, or has semantics still to be defined."""
    ws = load_workbook(WORKBOOK)[SHEET]
    headers = {section for section, _ in _LIST_FIELDS}
    section, where = None, {}
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row, 1).value
        if label in headers:
            section = label
        elif isinstance(label, str):
            where[f"B{row}"] = (section, label)
    checked = set()
    for dv in ws.data_validations.dataValidation:
        cell = str(dv.sqref)
        path = _LIST_FIELDS[where[cell]]
        options = [o.strip() for o in dv.formula1.strip('"').split(",")]
        for option in options:
            known = (option in _SUPPORTED[path]
                     or option in normalizer.UNSUPPORTED.get(path, {})
                     or option in normalizer.SEMANTICS_PENDING.get(path, {}))
            if path == "floating_leg.benchmark" and option in ("SARON", "INR OIS"):
                known = option in normalizer.UNSUPPORTED[path]
            assert known, f"{path}: workbook option {option!r} is not classified"
        checked.add(path)
    assert checked == set(_LIST_FIELDS.values())


def test_pending_options_remain_in_the_workbook():
    """Percentage per Period and the interpolation labels stay in the input contract."""
    ws = load_workbook(WORKBOOK)[SHEET]
    lists = " ".join(dv.formula1 for dv in ws.data_validations.dataValidation)
    for option in ("Percentage per Period", "Log Linear", "Linear", "Piecewise Linear (Cont)"):
        assert option in lists


# --- 9. every case worksheet (one worksheet = one test case) ---------------------------------

# Worksheets present and the case each one holds. Overnight-leg cases only: the term-rate
# cases (TC2, TC5, TC7, TC8, TC9) need reset inputs this contract does not have yet.
SHEET_CASES = {"TC001": "TC1", "TC003": "TC3", "TC004": "TC4", "TC006": "TC6", "TC010": "TC10"}


def test_workbook_holds_exactly_the_supported_cases():
    wb = load_workbook(WORKBOOK)
    assert wb.sheetnames == list(SHEET_CASES)
    template = sorted((str(d.sqref), d.formula1) for d in wb[SHEET].data_validations.dataValidation)
    for ws in wb.worksheets:  # same layout and drop-down lists on every sheet
        assert [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)] == [
            wb[SHEET].cell(r, 1).value for r in range(1, wb[SHEET].max_row + 1)]
        assert sorted((str(d.sqref), d.formula1)
                      for d in ws.data_validations.dataValidation) == template


@pytest.mark.parametrize("sheet, case", sorted(SHEET_CASES.items()))
def test_sheet_reproduces_the_existing_trade_file(sheet, case):
    trade_file = ROOT / "examples" / "irs" / f"{case.lower()}_trade.json"
    normalized = normalize(read_trade_sheet(WORKBOOK, sheet))
    assert normalized.trade_document == _json(trade_file)
    assert normalized.notes == ()


@pytest.mark.parametrize("sheet, case", sorted(SHEET_CASES.items()))
def test_sheet_prices_to_the_approved_baseline(sheet, case):
    baseline = _json(suite.expected_path(case))
    assert baseline["provenance"]["kind"] == "optima_regression_baseline"
    priced = price_workbook(WORKBOOK, sheet)
    assert priced.market_files == engine.resolve_market_files(
        _json(ROOT / "examples" / "irs" / f"{case.lower()}_trade.json")).files
    result = suite.CaseResult(case, "PASS")
    suite.compare(_suite_outputs(priced), baseline["expected"], result)
    assert result.failures == []
    assert result.fields_compared == ["npv", "fixed_leg_npv", "floating_leg_npv", "cashflows",
                                      "settled"]
