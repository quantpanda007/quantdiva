"""Read one trade worksheet of ``Optima_IRS_Test_Case_Inputs_v4.xlsx`` into the
source-independent trade inputs (``services.irs_inputs.contract``).

This is the only module that knows about Excel. It finds each field by its section
header and its label in column A (value in column B), so it does not depend on fixed cell
addresses, and it handles what is specific to Excel:

* date cells -> ``date``;
* percentage-formatted cells -> the percent number (a cell showing 5.0000% holds 0.05 and
  is read as 5). "Fixed Rate" must be percentage-formatted: a plain number there would be
  ambiguous (5 or 0.05?), so it is rejected;
* the formula cell "Snapshot Folder" -> its value as last calculated by Excel;
* the amortisation table -> one row per period up to the last filled row.

It does not validate or convert values beyond that (``normalizer`` does), and it applies no
defaults: an empty cell is ``None``. A workbook whose layout is not the expected one (a
section or label missing, an unknown label, data outside the table) is rejected.
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from services.irs_inputs.contract import SCHEDULE_ROW_KEYS, TradeInputs


class WorkbookReadError(ValueError):
    """The workbook or worksheet cannot be read as a trade input sheet."""


# section header (column A) -> (contract section, {label: (field, kind)})
# kind: "text" | "number" | "date" | "percent"
_SECTIONS: Dict[str, Tuple[str, Dict[str, Tuple[str, str]]]] = {
    "TRADE DETAILS": ("trade", {
        "Trade ID": ("trade_id", "text"),
        "Currency": ("currency", "text"),
        "Notional": ("notional", "number"),
        "Effective Date": ("effective_date", "date"),
        "Maturity Date": ("maturity_date", "date"),
        "Valuation Date": ("valuation_date", "date"),
    }),
    "SWAP LEGS — FIXED LEG": ("fixed_leg", {
        "Direction": ("direction", "text"),
        "Fixed Rate": ("rate_percent", "percent"),
        "Payment Frequency": ("payment_frequency", "text"),
        "Day Count": ("day_count", "text"),
        "Business Day Convention": ("business_day_convention", "text"),
        "Payment Calendar": ("calendar", "text"),
    }),
    "SWAP LEGS — FLOATING LEG": ("floating_leg", {
        "Direction": ("direction", "text"),
        "Benchmark": ("benchmark", "text"),
        "Index Tenor": ("index_tenor", "text"),
        "Spread": ("spread_bp", "number"),
        "Leverage": ("leverage", "number"),
        "Reset Frequency": ("reset_frequency", "text"),
        "Payment Frequency": ("payment_frequency", "text"),
        "Day Count": ("day_count", "text"),
        "Business Day Convention": ("business_day_convention", "text"),
        "Calendar": ("calendar", "text"),
        "Rate Timing": ("rate_timing", "text"),
        "Lookback": ("lookback_days", "number"),
        "Lookback Type": ("lookback_type", "text"),
        "Rate Cutoff": ("rate_cutoff_days", "number"),
        "Compounding": ("compounding", "text"),
        "Spread Treatment": ("spread_treatment", "text"),
    }),
    "NOTIONAL & AMORTISATION": ("notional", {
        "Principal Exchange": ("principal_exchange", "text"),
        "Amortisation Type": ("amortisation_type", "text"),
    }),
    "MARKET DATA REFERENCE": ("market_data", {
        "Curve": ("curve", "text"),
        "Snapshot Folder": ("snapshot_folder", "text"),
        "Curve Source File": ("curve_source_file", "text"),
        "Fixings Source File": ("fixings_source_file", "text"),
        "Reference Date": ("reference_date", "date"),
        "Interpolation": ("interpolation", "text"),
    }),
}
_SCHEDULE_HEADER = "Amortisation Schedule"
_SCHEDULE_COLUMNS = ("Period", "Effective Date", "End Date", "Payment Date", "Notional",
                     "Amortisation (%)", "Amortisation Amount", "Action")
_SCHEDULE_KINDS = ("number", "date", "date", "date", "number", "percent", "number", "text")


def _cell_value(cell, kind: str, where: str) -> Any:
    value = cell.value
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, str) and value.startswith("#"):
        raise WorkbookReadError(f"{where}: cell shows the error {value!r}")
    if kind == "date":
        return value.date() if isinstance(value, datetime) else value
    if kind == "percent":
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            if "%" not in (cell.number_format or ""):
                raise WorkbookReadError(f"{where}: must be a percentage-formatted cell "
                                        f"(e.g. 5.00%), got the plain number {value!r}")
            return Decimal(str(value)) * 100
        return value  # text: left to the normalizer to reject
    if isinstance(value, str):
        return value.strip()
    return value


def _label(cell) -> Optional[str]:
    value = cell.value
    return value.strip() if isinstance(value, str) and value.strip() else None


def read_trade_sheet(path, sheet: Optional[str] = None) -> TradeInputs:
    """Trade inputs from one worksheet. ``sheet`` may be omitted when the workbook has a
    single worksheet."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment problem
        raise WorkbookReadError("reading Excel inputs needs the openpyxl package") from exc
    path = Path(path)
    try:
        workbook = load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001 - any unreadable file is reported the same way
        raise WorkbookReadError(f"{path.name}: cannot open as an Excel workbook ({exc})") from exc
    try:
        if sheet is None:
            if len(workbook.sheetnames) != 1:
                raise WorkbookReadError(f"{path.name}: has {len(workbook.sheetnames)} "
                                        f"worksheets {workbook.sheetnames}; name the one to read")
            sheet = workbook.sheetnames[0]
        if sheet not in workbook.sheetnames:
            raise WorkbookReadError(f"{path.name}: no worksheet {sheet!r}; worksheets: "
                                    f"{workbook.sheetnames}")
        return _read(workbook[sheet], f"{path.name}[{sheet}]")
    finally:
        workbook.close()


def _read(ws, source: str) -> Dict[str, Any]:
    result: Dict[str, Dict[str, Any]] = {name: {} for name, _ in _SECTIONS.values()}
    seen_sections = set()
    current: Optional[Tuple[str, Dict[str, Tuple[str, str]]]] = None
    schedule_rows: Optional[List[Dict[str, Any]]] = None
    schedule_header_row: Optional[int] = None
    blank_after_data = False

    for row in range(1, ws.max_row + 1):
        label = _label(ws.cell(row, 1))
        if label in _SECTIONS:
            if label in seen_sections:
                raise WorkbookReadError(f"{source}: section {label!r} appears twice")
            seen_sections.add(label)
            current, schedule_rows = _SECTIONS[label], None
            continue
        if label == _SCHEDULE_HEADER:
            if current is None or current[0] != "notional":
                raise WorkbookReadError(f"{source} row {row}: {label!r} outside NOTIONAL & "
                                        "AMORTISATION")
            if "schedule" in result["notional"]:
                raise WorkbookReadError(f"{source} row {row}: second amortisation table")
            schedule_rows, schedule_header_row = [], None
            result["notional"]["schedule"] = schedule_rows
            continue
        if schedule_rows is not None:
            if schedule_header_row is None:
                header = tuple(_label(ws.cell(row, c)) for c in range(1, 9))
                if header == (None,) * 8:
                    continue
                if header != _SCHEDULE_COLUMNS:
                    raise WorkbookReadError(f"{source} row {row}: amortisation table header "
                                            f"is {header}, expected {_SCHEDULE_COLUMNS}")
                schedule_header_row = row
                continue
            values = [_cell_value(ws.cell(row, c), kind, f"{source} {ws.cell(row, c).coordinate}")
                      for c, kind in zip(range(1, 9), _SCHEDULE_KINDS)]
            if all(v is None for v in values[1:]):  # template row: period number only
                if schedule_rows:
                    blank_after_data = True
                continue
            if blank_after_data:
                raise WorkbookReadError(f"{source} row {row}: amortisation row after an empty "
                                        "row; periods must be consecutive")
            schedule_rows.append(dict(zip(SCHEDULE_ROW_KEYS, values)))
            continue
        if label is None:
            if current and any(ws.cell(row, c).value not in (None, "") for c in (2,)):
                raise WorkbookReadError(f"{source} B{row}: value without a label in column A")
            continue
        if current is None:
            continue  # title rows above the first section
        section, fields = current
        if label not in fields:
            raise WorkbookReadError(f"{source} A{row}: unknown label {label!r} in section "
                                    f"{section}; expected {sorted(fields)}")
        name, kind = fields[label]
        if name in result[section]:
            raise WorkbookReadError(f"{source} A{row}: label {label!r} appears twice")
        result[section][name] = _cell_value(ws.cell(row, 2), kind, f"{source} {label} (B{row})")

    missing_sections = [h for h in _SECTIONS if h not in seen_sections]
    if missing_sections:
        raise WorkbookReadError(f"{source}: sections not found: {missing_sections}")
    for header, (section, fields) in _SECTIONS.items():
        absent = [label for label, (name, _) in fields.items() if name not in result[section]]
        if absent:
            raise WorkbookReadError(f"{source}: section {header!r} has no rows {absent}; this "
                                    "is not the v4 input layout")
    if schedule_header_row is None:
        raise WorkbookReadError(f"{source}: amortisation table not found")
    return result
