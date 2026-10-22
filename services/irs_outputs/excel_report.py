"""Write one priced IRS trade to an Excel results workbook.

Sheets:

* **Summary**: trade, swap NPV, fixed- and floating-leg NPVs, market data used, notes
  (e.g. the provisional TC2 observation method), sign convention.
* **Fixed Leg** and **Floating Leg**: one row per contractual period, with the notional,
  its amortisation from the previous period, day count fraction, rate, cashflow amount,
  discount factor and present value. Periods paid on or before the valuation date are listed
  as "Settled", with no amounts (they are excluded from the NPV).
* **Net Cashflow**: one row per payment date: outstanding notional, amortisation, fixed
  amount, floating amount, net amount, discount factor and net present value.

Every number is the engine's own output (``services.irs_inputs.pricing.PricedTrade``); the
writer computes only the per-date net (fixed + floating) and the notional change between
periods, from the trade's contractual periods. Totals are the engine's NPVs. Amounts are
signed from the trade's point of view: the paid leg is negative.

The workbook is plain values (no formulas), so it reads the same in any tool.
"""

from __future__ import annotations

from collections import OrderedDict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from services.irs_inputs.pricing import PricedTrade


class ReportExistsError(FileExistsError):
    """The output file exists; it is never overwritten unless asked."""


_HEADER_FILL = "1F4E78"
_TOTAL_FILL = "D9EAF7"
_FORMATS = {"date": "yyyy-mm-dd", "money": "#,##0.00", "rate": "0.000000%", "df": "0.0000000",
            "dcf": "0.000000", "bp": "0.00", "int": "0", "text": "@"}


def _periods(trade, leg: str):
    return trade.fixed_periods if leg == "fixed" else trade.floating_periods


def _leg_rows(priced: PricedTrade, leg: str) -> List[Dict[str, Any]]:
    """Every contractual period of a leg: priced rows plus settled periods."""
    instrument = priced.engine_outputs[0]
    priced_rows = {r.period: r for r in priced.cashflows if r.leg == leg}
    settled = {s.period: s for s in priced.settled if s.leg == leg}
    rows, previous = [], None
    for number, period in enumerate(_periods(instrument.trade, leg), start=1):
        notional = float(period.notional.amount)
        amortisation = 0.0 if previous is None else previous - notional
        previous = notional
        row = priced_rows.get(number)
        base = {"period": number, "accrual_start": period.accrual_start,
                "accrual_end": period.accrual_end, "payment_date": period.payment_date,
                "notional": notional, "amortisation": amortisation}
        if row is not None:
            if (row.accrual_start, row.accrual_end, row.payment_date) != (
                    period.accrual_start, period.accrual_end, period.payment_date):
                raise ValueError(f"{leg} period {number}: engine row dates differ from the trade")
            base.update(status="Projected", side=row.side, dcf=row.day_count_fraction,
                        rate=row.rate, spread=row.spread, amount=row.amount,
                        df=row.discount_factor, pv=row.present_value)
        elif number in settled:
            base.update(status="Settled", side=settled[number].side, dcf=None, rate=None,
                        spread=None, amount=None, df=None, pv=None)
        else:
            raise ValueError(f"{leg} period {number}: neither priced nor settled")
        rows.append(base)
    return rows


def _net_rows(fixed: Sequence[Dict[str, Any]], floating: Sequence[Dict[str, Any]]):
    by_date: "OrderedDict[Any, Dict[str, Any]]" = OrderedDict()
    for leg, rows in (("fixed", fixed), ("floating", floating)):
        for r in rows:
            entry = by_date.setdefault(r["payment_date"], {
                "payment_date": r["payment_date"], "status": r["status"],
                "notional": r["notional"], "amortisation": r["amortisation"],
                "fixed": None, "floating": None, "df": None})
            if entry["status"] != r["status"]:
                raise ValueError(f"{r['payment_date']}: legs disagree on settlement")
            entry[leg] = r["amount"]
            if r["df"] is not None:
                if entry["df"] is not None and abs(entry["df"] - r["df"]) > 1e-15:
                    raise ValueError(f"{r['payment_date']}: legs use different discount factors")
                entry["df"] = r["df"]
            entry[f"{leg}_pv"] = r["pv"]
    rows = []
    for entry in sorted(by_date.values(), key=lambda e: e["payment_date"]):
        if entry["status"] == "Projected":
            entry["net"] = (entry["fixed"] or 0.0) + (entry["floating"] or 0.0)
            entry["net_pv"] = (entry.get("fixed_pv") or 0.0) + (entry.get("floating_pv") or 0.0)
        else:
            entry["net"] = entry["net_pv"] = None
        rows.append(entry)
    return rows


def _table(ws, start_row: int, headers: Sequence[str], rows: Sequence[Sequence[Any]],
           formats: Sequence[str], total: Optional[Sequence[Any]] = None) -> int:
    from openpyxl.styles import Alignment, Font, PatternFill

    for col, title in enumerate(headers, start=1):
        cell = ws.cell(start_row, col, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r = start_row
    for values in rows:
        r += 1
        for col, (value, fmt) in enumerate(zip(values, formats), start=1):
            cell = ws.cell(r, col, value)
            cell.number_format = _FORMATS[fmt]
    if total is not None:
        r += 1
        for col, (value, fmt) in enumerate(zip(total, formats), start=1):
            cell = ws.cell(r, col, value)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=_TOTAL_FILL)
            if value is not None and not isinstance(value, str):
                cell.number_format = _FORMATS[fmt]
    for col, title in enumerate(headers, start=1):
        letter = ws.cell(start_row, col).column_letter
        ws.column_dimensions[letter].width = max(12, min(22, len(title) + 4))
    ws.freeze_panes = ws.cell(start_row + 1, 1)
    return r


def _leg_sheet(wb, title: str, leg: str, rows, leg_npv: float) -> None:
    ws = wb.create_sheet(title)
    floating = leg == "floating"
    headers = ["Period", "Status", "Side", "Accrual Start", "Accrual End", "Payment Date",
               "Notional", "Amortisation", "Day Count Fraction",
               "Rate incl. Spread" if floating else "Fixed Rate"]
    formats = ["int", "text", "text", "date", "date", "date", "money", "money", "dcf", "rate"]
    if floating:
        headers.append("Spread (bp)")
        formats.append("bp")
    headers += ["Amount", "Discount Factor", "Present Value"]
    formats += ["money", "df", "money"]
    data = []
    for r in rows:
        values = [r["period"], r["status"], r["side"].capitalize(), r["accrual_start"],
                  r["accrual_end"], r["payment_date"], r["notional"], r["amortisation"],
                  r["dcf"], r["rate"]]
        if floating:
            values.append(None if r["spread"] is None else r["spread"] * 10000)
        values += [r["amount"], r["df"], r["pv"]]
        data.append(values)
    total = [None] * len(headers)
    total[1], total[-1] = f"{title} NPV", leg_npv
    _table(ws, 1, headers, data, formats, total)


def write_pricing_report(priced: PricedTrade, path, *, source: str = "",
                         overwrite: bool = False) -> Path:
    """Write ``priced`` to ``path`` (.xlsx). Refuses an existing file unless ``overwrite``."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"{path}: the results file must be .xlsx")
    if path.exists() and not overwrite:
        raise ReportExistsError(f"{path} exists; not overwritten (use overwrite)")
    instrument = priced.engine_outputs[0]
    trade = instrument.trade
    fixed, floating = _leg_rows(priced, "fixed"), _leg_rows(priced, "floating")
    net = _net_rows(fixed, floating)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "OPTIMA | IRS PRICING RESULT"
    ws["A1"].font = Font(bold=True, size=14)
    terms = trade.floating_terms
    summary = [
        ("Trade ID", priced.trade_id, "text"),
        ("Currency", priced.currency, "text"),
        ("Valuation Date", priced.valuation_date, "date"),
        ("Input", source or "-", "text"),
        ("Fixed Leg", trade.fixed_leg_side.capitalize(), "text"),
        ("Floating Benchmark", f"{terms.benchmark} {terms.benchmark_tenor}", "text"),
        ("Swap NPV", priced.npv, "money"),
        ("Fixed Leg NPV", priced.fixed_leg_npv, "money"),
        ("Floating Leg NPV", priced.floating_leg_npv, "money"),
        ("Pricing", f"{priced.engine_outputs[2].engine_used} engine, "
                    f"{priced.engine_outputs[2].model_used} model", "text"),
        ("Market Data", "\n".join(p.name for p in priced.market_files), "text"),
        ("Settled Periods", f"{len(priced.settled)} (paid on or before the valuation date; "
                            "excluded from the NPV)", "text"),
        ("Notes", "\n".join(priced.notes) or "-", "text"),
        ("Sign Convention", "Amounts and present values are signed from the trade's point of "
                            "view: the paid leg is negative.", "text"),
        ("Generated (UTC)", datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None),
         "datetime"),
    ]
    for i, (label, value, fmt) in enumerate(summary, start=3):
        ws.cell(i, 1, label).font = Font(bold=True)
        cell = ws.cell(i, 2, value)
        cell.number_format = "yyyy-mm-dd hh:mm:ss" if fmt == "datetime" else _FORMATS[fmt]
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        if label == "Swap NPV":
            cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 70

    _leg_sheet(wb, "Fixed Leg", "fixed", fixed, priced.fixed_leg_npv)
    _leg_sheet(wb, "Floating Leg", "floating", floating, priced.floating_leg_npv)

    ws = wb.create_sheet("Net Cashflow")
    headers = ["Payment Date", "Status", "Notional", "Amortisation", "Fixed Amount",
               "Floating Amount", "Net Amount", "Discount Factor", "Net Present Value"]
    formats = ["date", "text", "money", "money", "money", "money", "money", "df", "money"]
    data = [[e["payment_date"], e["status"], e["notional"], e["amortisation"], e["fixed"],
             e["floating"], e["net"], e["df"], e["net_pv"]] for e in net]
    total = [None] * len(headers)
    total[1], total[-1] = "Swap NPV", priced.npv
    _table(ws, 1, headers, data, formats, total)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
